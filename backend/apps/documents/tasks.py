from celery import shared_task
from django.utils import timezone
from google.cloud import vision
import PyPDF2
import docx
import openpyxl
import io
import boto3
from django.conf import settings
from apps.accounts.middleware import TenantSchemaContext
from .models import Document, Photo, ProcessingJob


@shared_task(bind=True)
def process_document(self, document_id, tenant_schema):
    """
    Process document: extract text, parse content, create embeddings
    """
    import logging
    logger = logging.getLogger(__name__)
    
    document = None
    with TenantSchemaContext(tenant_schema):
        try:
            document = Document.objects.get(id=document_id)
            document.processing_status = 'processing'
            document.save()

            # Download file from S3 or local storage
            try:
                file_content = _download_file(document.file_path)
                if not file_content:
                    raise ValueError("File content is empty")
            except Exception as e:
                raise ValueError(f"Failed to download file: {str(e)}")

            # Extract text based on file type
            text = ""
            try:
                if document.file_type == 'pdf':
                    text = _extract_text_from_pdf(file_content)
                elif document.file_type == 'docx':
                    text = _extract_text_from_docx(file_content)
                elif document.file_type == 'txt':
                    text = file_content.decode('utf-8')
                elif document.file_type == 'csv':
                    text = _extract_text_from_csv(file_content)
                elif document.file_type in ['xlsx', 'xls']:
                    text = _extract_text_from_excel(file_content)
                else:
                    logger.warning(f"Unknown file type: {document.file_type}")
                    text = ""
            except Exception as e:
                logger.error(f"Error extracting text from {document.file_type}: {e}")
                text = f"[Error extracting text: {str(e)}]"

            # Save extracted text (even if empty or error)
            document.extracted_text = text
            document.is_processed = True
            document.processing_status = 'completed'
            document.processed_at = timezone.now()
            document.save()

            # Create embeddings only if we have text
            if text and text.strip():
                from apps.embeddings.tasks import create_embeddings
                create_embeddings.delay(
                    content=text,
                    source_type='document',
                    source_id=document.id,
                    tenant_schema=tenant_schema
                )
            else:
                logger.warning(f"Document {document_id} has no extractable text, skipping embeddings")

            return f"Document {document_id} processed successfully"

        except Exception as e:
            logger.error(f"Error processing document {document_id}: {e}", exc_info=True)
            if document:
                try:
                    document.processing_status = 'failed'
                    document.processing_error = str(e)
                    document.save()
                except:
                    pass
            raise


@shared_task(bind=True)
def process_photo(self, photo_id, tenant_schema):
    """
    Process photo: Google Vision API analysis, OCR, create embeddings
    """
    import logging
    from google.api_core import exceptions as google_exceptions
    logger = logging.getLogger(__name__)
    
    photo = None
    with TenantSchemaContext(tenant_schema):
        try:
            photo = Photo.objects.get(id=photo_id)
            photo.processing_status = 'processing'
            photo.save()

            # Validate file size (max 20MB for Google Vision API)
            if photo.file_size > 20 * 1024 * 1024:
                raise ValueError(f"File too large: {photo.file_size / (1024*1024):.1f}MB. Maximum is 20MB.")

            # Download image from S3 or local storage
            try:
                image_content = _download_file(photo.file_path)
                if not image_content:
                    raise ValueError("File content is empty")
                if len(image_content) == 0:
                    raise ValueError("Downloaded file is empty")
            except Exception as e:
                raise ValueError(f"Failed to download file: {str(e)}")

            # Initialize Google Vision client
            try:
                client = vision.ImageAnnotatorClient()
                image = vision.Image(content=image_content)
            except Exception as e:
                raise ValueError(f"Failed to initialize Vision API client: {str(e)}")

            # Initialize results with defaults
            labels = []
            text = ""
            objects = []
            faces = []
            colors = []

            # Perform multiple detections with individual error handling
            # 1. Label detection
            try:
                labels_response = client.label_detection(image=image)
                if labels_response and hasattr(labels_response, 'label_annotations'):
                    labels = [
                        {
                            'description': label.description,
                            'score': float(label.score)
                        }
                        for label in labels_response.label_annotations
                    ]
                    logger.debug(f"Detected {len(labels)} labels for photo {photo_id}")
            except google_exceptions.GoogleAPIError as e:
                logger.error(f"Google Vision API error in label detection: {e}")
            except Exception as e:
                logger.warning(f"Error in label detection: {e}")

            # 2. Text detection (OCR)
            try:
                text_response = client.text_detection(image=image)
                if text_response and hasattr(text_response, 'text_annotations'):
                    if text_response.text_annotations and len(text_response.text_annotations) > 0:
                        text = text_response.text_annotations[0].description or ""
                        logger.debug(f"Detected {len(text)} characters of text for photo {photo_id}")
            except google_exceptions.GoogleAPIError as e:
                logger.error(f"Google Vision API error in text detection: {e}")
            except Exception as e:
                logger.warning(f"Error in text detection: {e}")

            # 3. Object detection
            try:
                objects_response = client.object_localization(image=image)
                if objects_response and hasattr(objects_response, 'localized_object_annotations'):
                    objects = [
                        {
                            'name': obj.name,
                            'score': float(obj.score)
                        }
                        for obj in objects_response.localized_object_annotations
                    ]
                    logger.debug(f"Detected {len(objects)} objects for photo {photo_id}")
            except google_exceptions.GoogleAPIError as e:
                logger.error(f"Google Vision API error in object detection: {e}")
            except Exception as e:
                logger.warning(f"Error in object detection: {e}")

            # 4. Face detection
            try:
                faces_response = client.face_detection(image=image)
                if faces_response and hasattr(faces_response, 'face_annotations'):
                    faces = [
                        {
                            'joy': face.joy_likelihood.name if hasattr(face.joy_likelihood, 'name') else str(face.joy_likelihood),
                            'sorrow': face.sorrow_likelihood.name if hasattr(face.sorrow_likelihood, 'name') else str(face.sorrow_likelihood),
                            'anger': face.anger_likelihood.name if hasattr(face.anger_likelihood, 'name') else str(face.anger_likelihood)
                        }
                        for face in faces_response.face_annotations
                    ]
                    logger.debug(f"Detected {len(faces)} faces for photo {photo_id}")
            except google_exceptions.GoogleAPIError as e:
                logger.error(f"Google Vision API error in face detection: {e}")
            except Exception as e:
                logger.warning(f"Error in face detection: {e}")

            # 5. Color detection
            try:
                colors_response = client.image_properties(image=image)
                if colors_response and hasattr(colors_response, 'image_properties_annotation'):
                    if (colors_response.image_properties_annotation and 
                        hasattr(colors_response.image_properties_annotation, 'dominant_colors') and
                        colors_response.image_properties_annotation.dominant_colors):
                        colors = [
                            {
                                'color': {
                                    'red': int(color.color.red),
                                    'green': int(color.color.green),
                                    'blue': int(color.color.blue)
                                },
                                'score': float(color.score),
                                'pixel_fraction': float(color.pixel_fraction)
                            }
                            for color in colors_response.image_properties_annotation.dominant_colors.colors
                        ]
                        logger.debug(f"Detected {len(colors)} dominant colors for photo {photo_id}")
            except google_exceptions.GoogleAPIError as e:
                logger.error(f"Google Vision API error in color detection: {e}")
            except Exception as e:
                logger.warning(f"Error in color detection: {e}")

            # Save results (even if some detections failed)
            photo.labels = labels
            photo.text = text
            photo.detected_objects = objects
            photo.faces = faces
            photo.colors = colors
            photo.is_processed = True
            photo.processing_status = 'completed'
            photo.processed_at = timezone.now()
            photo.save()

            logger.info(f"Photo {photo_id} processed: {len(labels)} labels, {len(text)} chars text, {len(objects)} objects, {len(faces)} faces, {len(colors)} colors")

            # Create embeddings from extracted text and labels
            # Безпечно обробляємо labels (може бути None або порожнім списком)
            labels_text = ""
            if labels and isinstance(labels, list):
                labels_text = ' '.join([l.get('description', '') for l in labels if isinstance(l, dict) and l.get('description')])
            
            combined_text = f"{text} {labels_text}".strip()
            
            # Створюємо ембедінги тільки якщо є контент
            if combined_text:
                from apps.embeddings.tasks import create_embeddings
                create_embeddings.delay(
                    content=combined_text,
                    source_type='photo',
                    source_id=photo.id,
                    tenant_schema=tenant_schema
                )
                logger.info(f"Created embedding task for photo {photo_id}")
            else:
                logger.warning(f"Photo {photo_id} has no extractable text or labels, skipping embeddings")

            return f"Photo {photo_id} processed successfully"

        except Exception as e:
            logger.error(f"Error processing photo {photo_id}: {e}", exc_info=True)
            if photo:
                try:
                    photo.processing_status = 'failed'
                    photo.processing_error = str(e)
                    photo.save()
                except:
                    pass
            raise


@shared_task
def cleanup_old_files():
    """
    Delete old files from storage to save space
    """
    # TODO: Implement cleanup logic based on organization settings
    pass


# Helper functions

def _download_file(file_path):
    """Download file from S3 or local storage"""
    from django.core.files.storage import default_storage
    
    if settings.USE_S3:
        s3 = boto3.client('s3')
        response = s3.get_object(
            Bucket=settings.AWS_STORAGE_BUCKET_NAME,
            Key=file_path
        )
        return response['Body'].read()
    else:
        # Use default_storage to handle relative paths correctly
        if default_storage.exists(file_path):
            with default_storage.open(file_path, 'rb') as f:
                return f.read()
        else:
            # Fallback to direct file path (absolute)
            import os
            if os.path.exists(file_path):
                with open(file_path, 'rb') as f:
                    return f.read()
            else:
                raise FileNotFoundError(f"File not found: {file_path}")


def _extract_text_from_pdf(file_content):
    """Extract text from PDF"""
    import logging
    logger = logging.getLogger(__name__)
    
    try:
        pdf_reader = PyPDF2.PdfReader(io.BytesIO(file_content))
        
        # Check if PDF is encrypted
        if pdf_reader.is_encrypted:
            logger.warning("PDF is encrypted, attempting to decrypt with empty password")
            try:
                pdf_reader.decrypt("")
            except:
                raise ValueError("PDF is encrypted and cannot be decrypted")
        
        text = ""
        for i, page in enumerate(pdf_reader.pages):
            try:
                page_text = page.extract_text()
                if page_text:
                    text += page_text + "\n"
            except Exception as e:
                logger.warning(f"Error extracting text from page {i+1}: {e}")
                continue
        
        if not text.strip():
            logger.warning("No text extracted from PDF (may be image-based)")
            return "[PDF contains no extractable text - may be image-based]"
        
        return text
    except Exception as e:
        logger.error(f"Error extracting text from PDF: {e}")
        raise ValueError(f"Failed to extract text from PDF: {str(e)}")


def _extract_text_from_docx(file_content):
    """Extract text from DOCX"""
    doc = docx.Document(io.BytesIO(file_content))
    text = "\n".join([paragraph.text for paragraph in doc.paragraphs])
    return text


def _extract_text_from_csv(file_content):
    """Extract text from CSV"""
    import csv
    import logging
    logger = logging.getLogger(__name__)
    
    try:
        # Try to decode as UTF-8
        try:
            content_str = file_content.decode('utf-8')
        except UnicodeDecodeError:
            # Try other encodings
            try:
                content_str = file_content.decode('latin-1')
            except:
                content_str = file_content.decode('utf-8', errors='ignore')
        
        csv_reader = csv.reader(io.StringIO(content_str))
        text = ""
        for row in csv_reader:
            row_text = " ".join([str(cell) for cell in row if cell])
            if row_text.strip():
                text += row_text + "\n"
        return text
    except Exception as e:
        logger.error(f"Error extracting text from CSV: {e}")
        raise ValueError(f"Failed to extract text from CSV: {str(e)}")


def _extract_text_from_excel(file_content):
    """Extract text from Excel"""
    import logging
    logger = logging.getLogger(__name__)
    
    try:
        workbook = openpyxl.load_workbook(io.BytesIO(file_content), data_only=True)
        text = ""
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows(values_only=True):
                row_text = " ".join([str(cell) for cell in row if cell is not None])
                if row_text.strip():
                    text += row_text + "\n"
        return text
    except Exception as e:
        logger.error(f"Error extracting text from Excel: {e}")
        raise ValueError(f"Failed to extract text from Excel: {str(e)}")
